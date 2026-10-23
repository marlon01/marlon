#!/usr/bin/env python3
"""Construit les manuels (docs/manuels/<module>/*.md) en .docx puis en .pdf.

Chaîne : pandoc (Markdown → Word, avec un modèle de styles) → mise en forme python-docx
(page de garde, en-tête, pied de page numéroté, tableaux) → LibreOffice (mise à jour du
sommaire, enregistrement .docx et export .pdf).

Usage : python3 tools/build_manuels.py [fichier.md ...]   (par défaut : tous les manuels)
Dépendances : pandoc, LibreOffice (soffice + module Python uno), python-docx, PyYAML.
"""
import copy
import os
import re
import subprocess
import sys
import tempfile
import time

import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MANUELS = os.path.join(ROOT, 'docs', 'manuels')
INK = RGBColor(0x1A, 0x23, 0x33)
MUTED = RGBColor(0x5B, 0x65, 0x75)
ACCENTS = {  # couleur d'accent par produit
    'norma_facture': '0B5FFF',
    'selecta_concours': '6D28D9',
    'selecta_saas': '6D28D9',
    'liputa_couture': 'B45309',
}
EDITEUR = "Naxoris — contact@naxoris.com — +243 828 317 540 · +243 831 509 420"


def rgb(hexa):
    return RGBColor.from_string(hexa)


# --------------------------------------------------------------------------- modèle de styles
def build_reference(path, accent):
    with open(path, 'wb') as fh:
        fh.write(subprocess.check_output(['pandoc', '--print-default-data-file', 'reference.docx']))
    doc = Document(path)
    st = doc.styles

    from docx.enum.style import WD_STYLE_TYPE
    by_name = {x.name: x for x in st}
    if 'Source Code' not in by_name:
        sc = st.add_style('Source Code', WD_STYLE_TYPE.PARAGRAPH)
        sc.base_style = by_name['Normal']
        sc.paragraph_format.space_after = Pt(0)
        sc.paragraph_format.line_spacing = 1.0
        by_name['Source Code'] = sc

    def font(style, size=None, bold=None, color=None, italic=None, name='Calibri'):
        s = by_name[style]
        s.font.name = name
        rpr = s.element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.append(rfonts)
        for att in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rfonts.set(qn(att), name)
        for att in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
            if rfonts.get(qn(att)) is not None:
                del rfonts.attrib[qn(att)]
        if size:
            s.font.size = Pt(size)
        if bold is not None:
            s.font.bold = bold
        if italic is not None:
            s.font.italic = italic
        if color is not None:
            s.font.color.rgb = color
        return s

    for name in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
        s = font(name, 10.5, color=INK)
        s.paragraph_format.line_spacing = 1.15
        s.paragraph_format.space_after = Pt(5 if name != 'Compact' else 2)
        s.paragraph_format.space_before = Pt(0)
    by_name['Body Text'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    by_name['First Paragraph'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    sizes = {1: 18, 2: 14, 3: 12, 4: 11}
    for lvl, size in sizes.items():
        s = font('Heading %d' % lvl, size, bold=True, color=rgb(accent) if lvl == 1 else INK)
        s.font.italic = False
        pf = s.paragraph_format
        pf.space_before = Pt(18 if lvl == 1 else 12)
        pf.space_after = Pt(8 if lvl == 1 else 4)
        pf.keep_with_next = True
        pf.page_break_before = (lvl == 1)
    # Filet sous les chapitres
    h1 = by_name['Heading 1'].element.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '8'), ('w:space', '4'), ('w:color', accent)):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    h1.append(bdr)

    font('Title', 30, bold=True, color=INK).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font('Subtitle', 15, color=rgb(accent)).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font('Author', 11, color=MUTED).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font('Date', 11, color=MUTED).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font('TOC Heading', 18, bold=True, color=rgb(accent))
    s = font('Image Caption', 9, italic=True, color=MUTED)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_after = Pt(10)
    by_name['Captioned Figure'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    by_name['Captioned Figure'].paragraph_format.keep_with_next = True
    font('Source Code', 8.5, name='Consolas', color=INK)
    font('Verbatim Char', 9, name='Consolas', color=INK)
    s = font('Block Text', 10, color=INK)
    s.paragraph_format.left_indent = Cm(0.4)
    s.paragraph_format.right_indent = Cm(0)
    s.paragraph_format.space_before = Pt(4)
    s.paragraph_format.space_after = Pt(6)
    ppr = s.element.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    for k, v in (('w:val', 'single'), ('w:sz', '18'), ('w:space', '8'), ('w:color', accent)):
        left.set(qn(k), v)
    bdr.append(left)
    ppr.append(bdr)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), 'F3F5F9')
    ppr.append(shd)
    # Code : fond gris clair
    ppr = by_name['Source Code'].element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), 'F3F4F6')
    ppr.append(shd)
    # Marges A4
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.2)
        sec.header_distance = sec.footer_distance = Cm(1.1)
    doc.save(path)


# --------------------------------------------------------------------------- mise en forme
def field_run(paragraph, instr):
    run = paragraph.add_run()
    for kind, text in (('begin', None), (None, instr), ('separate', None), (None, '1'), ('end', None)):
        if kind:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), kind)
        elif text == instr:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = ' %s ' % instr
        else:
            el = OxmlElement('w:t')
            el.text = text
        run._r.append(el)
    return run


def style_tables(doc, accent):
    for table in doc.tables:
        tbl = table._tbl
        tblpr = tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement('w:%s' % edge)
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), '4')
            el.set(qn('w:color'), 'C9CED6')
            borders.append(el)
        old = tblpr.find(qn('w:tblBorders'))
        if old is not None:
            tblpr.remove(old)
        tblpr.append(borders)
        mar = OxmlElement('w:tblCellMar')
        for edge, w in (('top', 40), ('bottom', 40), ('left', 90), ('right', 90)):
            el = OxmlElement('w:%s' % edge)
            el.set(qn('w:w'), str(w))
            el.set(qn('w:type'), 'dxa')
            mar.append(el)
        tblpr.append(mar)
        spacer = OxmlElement('w:p')
        sp_ppr = OxmlElement('w:pPr')
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:before'), '0')
        sp.set(qn('w:after'), '60')
        sp.set(qn('w:line'), '120')
        sp.set(qn('w:lineRule'), 'exact')
        sp_ppr.append(sp)
        spacer.append(sp_ppr)
        tbl.addnext(spacer)
        for i, row in enumerate(table.rows):
            if i == 0:
                trpr = row._tr.get_or_add_trPr()
                h = OxmlElement('w:tblHeader')
                trpr.append(h)
            for cell in row.cells:
                tcpr = cell._tc.get_or_add_tcPr()
                if i == 0 or i % 2 == 0:
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:val'), 'clear')
                    shd.set(qn('w:fill'), accent if i == 0 else 'F6F7F9')
                    tcpr.append(shd)
                for p in cell.paragraphs:
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if i == 0:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def finish(docx_path, meta, accent):
    doc = Document(docx_path)
    body = doc.element.body
    title = meta.get('title', '')
    # Page de garde : bandeau, informations, saut de page avant le sommaire
    paras = doc.paragraphs
    first = paras[0]
    band = first.insert_paragraph_before()
    band.paragraph_format.space_before = Pt(110)
    run = band.add_run((meta.get('produit') or title.split('—')[0]).strip().upper())
    run.font.size, run.font.bold, run.font.color.rgb = Pt(11), True, rgb(accent)
    last_meta = None
    for p in doc.paragraphs:
        if p.style.name in ('Title', 'Subtitle', 'Author', 'Date'):
            last_meta = p
    anchor = last_meta or first
    lines = []
    if meta.get('version'):
        lines.append('Version du logiciel : %s' % meta['version'])
    lines.append('Éditeur : %s' % EDITEUR)
    if meta.get('classification'):
        lines.append(meta['classification'])
    from docx.text.paragraph import Paragraph
    cur = anchor._p
    for text in lines + [None]:
        el = OxmlElement('w:p')
        cur.addnext(el)
        cur = el
        p = Paragraph(el, anchor._parent)
        if text is None:
            p.add_run().add_break(WD_BREAK.PAGE)
        else:
            p.paragraph_format.space_before = Pt(16 if text is lines[0] else 2)
            r = p.add_run(text)
            r.font.size, r.font.color.rgb = Pt(10), MUTED

    # En-tête et pied de page (pas sur la page de garde)
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.text = ''
    r = hp.add_run(title)
    r.font.size, r.font.color.rgb = Pt(8.5), MUTED
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = sec.footer.paragraphs[0]
    fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('%s — page ' % meta.get('editeur', 'Naxoris'))
    r.font.size, r.font.color.rgb = Pt(8.5), MUTED
    run = field_run(fp, 'PAGE')
    run.font.size = Pt(8.5)
    r = fp.add_run(' / ')
    r.font.size, r.font.color.rgb = Pt(8.5), MUTED
    run = field_run(fp, 'NUMPAGES')
    run.font.size = Pt(8.5)

    style_tables(doc, accent)
    # Les chapitres sautent déjà de page ; le premier après le sommaire aussi.
    doc.core_properties.title = title
    doc.core_properties.subject = meta.get('subtitle', '')
    doc.core_properties.author = meta.get('editeur', 'Naxoris')
    doc.core_properties.language = 'fr-FR'
    doc.save(docx_path)


# --------------------------------------------------------------------------- LibreOffice
UNO_SCRIPT = r'''
import sys, time, uno
from com.sun.star.beans import PropertyValue

def prop(n, v):
    p = PropertyValue(); p.Name = n; p.Value = v; return p

port, src, out_docx, out_pdf = sys.argv[1:5]
local = uno.getComponentContext()
resolver = local.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver', local)
for _ in range(60):
    try:
        ctx = resolver.resolve('uno:socket,host=localhost,port=%s;urp;StarOffice.ComponentContext' % port)
        break
    except Exception:
        time.sleep(0.5)
desktop = ctx.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop', ctx)
doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(src), '_blank', 0, (prop('Hidden', True),))
doc.getDocumentIndexes()
for _ in range(2):  # 2 passes : les numéros de page du sommaire se stabilisent
    idx = doc.getDocumentIndexes()
    for i in range(idx.getCount()):
        idx.getByIndex(i).update()
    doc.refresh()
doc.storeToURL(uno.systemPathToFileUrl(out_docx), (prop('FilterName', 'MS Word 2007 XML'),))
doc.storeToURL(uno.systemPathToFileUrl(out_pdf), (prop('FilterName', 'writer_pdf_Export'),
                prop('FilterData', uno.Any('[]com.sun.star.beans.PropertyValue', tuple([
                    prop('UseTaggedPDF', True), prop('ExportBookmarks', True), prop('Quality', 90)])))))
doc.close(True)
'''


class Office:
    def __init__(self):
        self.port = '2%04d' % (os.getpid() % 10000)
        self.profile = tempfile.mkdtemp(prefix='lo-profile-')
        self.proc = subprocess.Popen([
            'soffice', '--headless', '--invisible', '--nologo', '--norestore',
            '-env:UserInstallation=file://%s' % self.profile,
            '--accept=socket,host=localhost,port=%s;urp;' % self.port],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        self.script = os.path.join(self.profile, 'convert.py')
        with open(self.script, 'w') as fh:
            fh.write(UNO_SCRIPT)
        time.sleep(2)

    def convert(self, src, out_docx, out_pdf):
        subprocess.run(['python3', self.script, self.port, src, out_docx, out_pdf], check=True, timeout=600)

    def close(self):
        self.proc.terminate()
        self.proc.wait(timeout=30)


def front_matter(md_path):
    with open(md_path, encoding='utf-8') as fh:
        text = fh.read()
    m = re.match(r'^---\n(.*?)\n---\n', text, re.S)
    return yaml.safe_load(m.group(1)) if m else {}


def build(md_path, office, workdir):
    module = os.path.basename(os.path.dirname(md_path))
    accent = ACCENTS.get(module, '1F3A5F')
    meta = front_matter(md_path)
    ref = os.path.join(workdir, 'ref-%s.docx' % accent)
    if not os.path.exists(ref):
        build_reference(ref, accent)
    base = os.path.splitext(md_path)[0]
    tmp = os.path.join(workdir, os.path.basename(base) + '-%s.docx' % module)
    subprocess.run([
        'pandoc', md_path, '-f', 'markdown+implicit_figures+pipe_tables+yaml_metadata_block',
        '-o', tmp, '--reference-doc', ref, '--toc', '--toc-depth=2', '--number-sections',
        '--resource-path', ROOT, '-M', 'toc-title=Sommaire', '-M', 'lang=fr-FR',
        '-M', 'author=%s' % meta.get('editeur', 'Naxoris'),
    ], check=True, cwd=ROOT)
    finish(tmp, meta, accent)
    office.convert(tmp, base + '.docx', base + '.pdf')
    print('ok', os.path.relpath(base, ROOT) + '.{docx,pdf}')


def main():
    files = sys.argv[1:] or sorted(
        os.path.join(dp, f) for dp, _d, fs in os.walk(MANUELS) for f in fs
        if f.endswith('.md') and f != 'CONVENTIONS.md' and dp != MANUELS)
    workdir = tempfile.mkdtemp(prefix='manuels-')
    office = Office()
    try:
        for f in files:
            build(os.path.abspath(f), office, workdir)
    finally:
        office.close()


if __name__ == '__main__':
    main()
